# -*- coding: utf-8 -*-
"""TerraPilot — fouilles, engins, incertitudes et suivi pédagogique.

Tout le code et les ressources de présentation sont regroupés ici.
Le registre SQLite est conservé hors du dossier de l'application.
"""
from __future__ import annotations

import base64
import copy
from contextlib import closing
from datetime import date, datetime, timedelta, timezone
import hashlib
import html
import io
import json
import math
import os
from pathlib import Path
import sqlite3
import uuid
import warnings

import numpy as np
import pandas as pd
import plotly.graph_objects as go
from scipy import stats
from PIL import Image, ImageOps
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Flowable, KeepTogether

APP='TerraPilot'
GREEN='#183F35'; COPPER='#B65C38'; SAND='#F4F0E7'; BLUE='#387D8A'
NAV={'Vue d’ensemble':'dashboard','Chantiers':'location_on','Fouilles et engins':'construction',
     'Simulations':'query_stats','Suivi des travaux':'engineering','Guide et exports':'menu_book'}
SOILS=['Non caractérisé','Argileux observé','Sableux observé','Graveleux observé','Rocheux observé','Hétérogène / remblai']
LIMIT=15*1024*1024
LOG_COLUMNS=['date','volume_m3','heures_actives','heures_arret','carburant_l','depense_fcfa','note']


def esc(x):return html.escape(str(x))
def uid():return uuid.uuid4().hex
def stamp():return datetime.now(timezone.utc).isoformat(timespec='microseconds')
def fmt(x,n=0):return f'{float(x):,.{n}f}'.replace(',',' ').replace('.',',')
def money(x):return fmt(x)+' FCFA'
def num(v,label,lo=0,hi=1e12,integer=False):
    try:x=float(str(v).replace(',','.'))
    except (ValueError,TypeError):raise ValueError(label+' : nombre valide attendu.') from None
    if not math.isfinite(x) or not lo<=x<=hi or (integer and not x.is_integer()):
        raise ValueError(f'{label} : valeur {"entière " if integer else ""}entre {lo:g} et {hi:g} attendue.')
    return int(x) if integer else x
def txt(v,label,limit=120,required=False):
    x=str(v or '').strip()
    if len(x)>limit or (required and not x):raise ValueError(f'{label} : texte {"obligatoire ; " if required else ""}{limit} caractères maximum.')
    return x
def iso(v):
    try:return date.fromisoformat(str(v))
    except (ValueError,TypeError):raise ValueError('Date attendue au format AAAA-MM-JJ.') from None
def empty():return dict(schema=1,projects=[],logs=[])
def code(rows):
    existing={r['id'] for r in rows};i=1
    while f'TP-{i:04}' in existing:i+=1
    return f'TP-{i:04}'


def defaults():
    return dict(mode='Fouille rectangulaire',length=10.,width=2.,depth=1.5,count=1,
                clearance=0.,slope=0.,bulking=20.,hours_day=8.,setup=0.,haul=0.)
def default_machines():
    return [dict(id='M1',name='Pelle compacte',low=4.,rate=5.,high=6.,hourly=20000.,fuel=4.,fuel_price=1000.,included=True),
            dict(id='M2',name='Pelle sur chenilles',low=10.,rate=15.,high=20.,hourly=45000.,fuel=12.,fuel_price=1000.,included=False),
            dict(id='M3',name='Équipe manuelle',low=.6,rate=1.,high=1.4,hourly=5000.,fuel=0.,fuel_price=0.,included=True)]
def default_sim():return dict(distribution='Triangulaire',count=10000,seed=42,stop_probability=.2,stop_min=2.,stop_max=8.,standby=.25)
def new_project(name='Nouveau chantier',city='À renseigner'):
    return dict(id='',name=name,city=city,start=date.today().isoformat(),soil=SOILS[0],water=None,
                note='',photo='',budget=150000.,deadline=1.,layers=[dict(name='Couche observée',thickness=2.)],
                geometry=defaults(),machines=default_machines(),selected='M1',simulation=default_sim())


def validate_geometry(raw):
    if not isinstance(raw,dict):raise ValueError('Paramètres de fouille manquants.')
    mode=raw.get('mode')
    if mode not in ('Fouille rectangulaire','Semelle avec espace de travail'):raise ValueError('Mode de fouille non reconnu.')
    bounds={'length':(.01,1000),'width':(.01,1000),'depth':(.01,30),'count':(1,1000),
            'clearance':(0,10),'slope':(0,3),'bulking':(0,100),'hours_day':(.1,24),'setup':(0,1e10),'haul':(0,1e7)}
    return dict(mode=mode,**{k:num(raw.get(k),k,*b,integer=k=='count') for k,b in bounds.items()})
def geometry(raw):
    g=validate_geometry(raw);space=g['clearance'] if g['mode'].startswith('Semelle') else 0.
    l=g['length']+2*space;w=g['width']+2*space;h=g['depth'];m=g['slope']
    top_l=l+2*m*h;top_w=w+2*m*h
    unit=h/6*(l*w+4*(l+m*h)*(w+m*h)+top_l*top_w)
    volume=unit*g['count']
    if volume>1e7:raise ValueError('Volume limité à 10 millions de m³.')
    return dict(bottom_l=l,bottom_w=w,top_l=top_l,top_w=top_w,depth=h,unit=unit,
                volume=volume,loose=volume*(1+g['bulking']/100),surface=top_l*top_w*g['count'])
def validate_machine(raw):
    if not isinstance(raw,dict):raise ValueError('Fiche du moyen de travail invalide.')
    out=dict(id=txt(raw.get('id'),'ID engin',50,True),name=txt(raw.get('name'),'Nom engin',80,True))
    for k,b in {'low':(.01,10000),'rate':(.01,10000),'high':(.01,10000),'hourly':(0,1e8),
                'fuel':(0,1000),'fuel_price':(0,10000)}.items():out[k]=num(raw.get(k),k,*b)
    if not out['low']<=out['rate']<=out['high']:raise ValueError('Rendement : minimum ≤ habituel ≤ maximum requis.')
    if type(raw.get('included')) is not bool:raise ValueError('Préciser si le carburant est inclus dans le tarif.')
    out['included']=raw['included'];return out
def validate_sim(raw):
    if not isinstance(raw,dict) or raw.get('distribution') not in ('Triangulaire','Uniforme','Fixe'):raise ValueError('Loi de rendement non reconnue.')
    bounds={'count':(1000,100000),'seed':(0,2**32-1),'stop_probability':(0,1),'stop_min':(0,500),'stop_max':(0,500),'standby':(0,1)}
    out=dict(distribution=raw['distribution'],**{k:num(raw.get(k),k,*b,integer=k in ('count','seed')) for k,b in bounds.items()})
    if out['stop_min']>out['stop_max']:raise ValueError('Arrêts : minimum ≤ maximum requis.')
    return out
def estimate(p,machine=None):
    g=validate_geometry(p['geometry']);v=geometry(g)
    m=validate_machine(machine or next(m for m in p['machines'] if m['id']==p['selected']))
    hours=v['volume']/m['rate'];fuel=hours*m['fuel'];location=hours*m['hourly']
    fuel_cost=0. if m['included'] else fuel*m['fuel_price'];haul=v['loose']*g['haul']
    total=location+fuel_cost+haul+g['setup']
    if total>1e12:raise ValueError('Estimation supérieure à 1 000 milliards de FCFA.')
    return dict(**v,hours=hours,days=hours/g['hours_day'],fuel=fuel,location=location,
                fuel_cost=fuel_cost,haul=haul,setup=g['setup'],total=total,machine=m['name'])


def photo_bytes(raw):
    if not raw:return ''
    if len(raw)>5*1024*1024:raise ValueError('Photo limitée à 5 Mo.')
    try:
        with warnings.catch_warnings():
            warnings.simplefilter('error',Image.DecompressionBombWarning)
            with Image.open(io.BytesIO(raw)) as image:
                if image.width*image.height>20_000_000:raise ValueError('Photo limitée à 20 mégapixels.')
                image=ImageOps.exif_transpose(image).convert('RGB');image.thumbnail((1200,800));b=io.BytesIO();image.save(b,'JPEG',quality=78)
                return base64.b64encode(b.getvalue()).decode()
    except (OSError,Image.DecompressionBombError,Image.DecompressionBombWarning):raise ValueError('Image JPEG, PNG ou WebP valide attendue.') from None
def validate_photo(value):
    if not isinstance(value,str) or len(value)>950000:raise ValueError('Photo trop volumineuse.')
    if value:
        try:
            raw=base64.b64decode(value,validate=True)
            with Image.open(io.BytesIO(raw)) as im:
                if im.format!='JPEG' or im.width*im.height>20_000_000:raise ValueError('Photo JPEG normalisée attendue.')
                im.verify()
        except Exception:raise ValueError('Photo de sauvegarde invalide.') from None
    return value
def validate_project(raw):
    if not isinstance(raw,dict):raise ValueError('Fiche de chantier invalide.')
    p=dict(id=txt(raw.get('id'),'ID chantier',50,True),name=txt(raw.get('name'),'Chantier',100,True),city=txt(raw.get('city'),'Localisation',100,True),
           start=iso(raw.get('start')).isoformat(),note=txt(raw.get('note'),'Observations',2000),soil=raw.get('soil'),photo=validate_photo(raw.get('photo','')),
           budget=num(raw.get('budget'),'Budget',0,1e12),deadline=num(raw.get('deadline'),'Échéance',.01,10000),
           water=None if raw.get('water') is None else num(raw['water'],'Profondeur eau',0,1000))
    if p['soil'] not in SOILS:raise ValueError('Description de sol non reconnue.')
    layers=raw.get('layers')
    if not isinstance(layers,list) or not 1<=len(layers)<=10:raise ValueError('Renseigner de 1 à 10 couches observées.')
    if not all(isinstance(r,dict) for r in layers):raise ValueError('Chaque couche doit avoir un nom et une épaisseur.')
    p['layers']=[dict(name=txt(r.get('name'),'Couche',80,True),thickness=num(r.get('thickness'),'Épaisseur de couche',.01,1000)) for r in layers]
    p['geometry']=validate_geometry(raw.get('geometry'));machines=raw.get('machines')
    if not isinstance(machines,list) or not 1<=len(machines)<=8:raise ValueError('Renseigner de 1 à 8 moyens de travail.')
    p['machines']=[validate_machine(m) for m in machines]
    if len({m['id'] for m in p['machines']})!=len(p['machines']):raise ValueError('Identifiants d’engins dupliqués.')
    p['selected']=raw.get('selected')
    if p['selected'] not in [m['id'] for m in p['machines']]:raise ValueError('Choisir un engin existant.')
    p['simulation']=validate_sim(raw.get('simulation'));estimate(p)
    return p
def validate_log(raw,p):
    if not isinstance(raw,dict):raise ValueError('Relevé invalide.')
    out=dict(id=txt(raw.get('id'),'ID relevé',50,True),project=p['id'],date=iso(raw.get('date')).isoformat(),
             note=txt(raw.get('note'),'Note',1000),engine=txt(raw.get('engine'),'Moyen utilisé',80,True))
    if not iso(p['start'])<=iso(out['date'])<=date.today():raise ValueError('La date du relevé doit être entre le début du chantier et aujourd’hui.')
    for key,hi in [('volume_m3',1e7),('heures_actives',24),('heures_arret',24),('carburant_l',1e6),('depense_fcfa',1e12)]:out[key]=num(raw.get(key),key,0,hi)
    if out['heures_actives']+out['heures_arret']>24:raise ValueError('Un relevé ne peut pas dépasser 24 heures au total.')
    if out['volume_m3']>0 and out['heures_actives']==0:raise ValueError('Un volume réalisé exige un temps actif positif.')
    if not any(out[k]>0 for k in ['volume_m3','heures_actives','heures_arret','carburant_l','depense_fcfa']):raise ValueError('Renseigner au moins une quantité mesurée.')
    return out
def validate_store(raw):
    if not isinstance(raw,dict) or raw.get('schema')!=1:raise ValueError('Sauvegarde TerraPilot version 1 attendue.')
    if not isinstance(raw.get('projects'),list) or len(raw['projects'])>30:raise ValueError('Limite de 30 chantiers.')
    if not isinstance(raw.get('logs'),list) or len(raw['logs'])>3000:raise ValueError('Limite de 3 000 relevés.')
    projects=[validate_project(p) for p in raw['projects']];lookup={p['id']:p for p in projects}
    if len(lookup)!=len(projects):raise ValueError('Identifiants de chantiers dupliqués.')
    logs=[]
    for row in raw['logs']:
        if not isinstance(row,dict) or row.get('project') not in lookup:raise ValueError('Relevé rattaché à un chantier inconnu.')
        logs.append(validate_log(row,lookup[row['project']]))
    if len({r['id'] for r in logs})!=len(logs):raise ValueError('Identifiants de relevés dupliqués.')
    data=dict(schema=1,projects=projects,logs=logs)
    if len(json.dumps(data).encode())>LIMIT:raise ValueError('Registre limité à 15 Mo.')
    return data
def records(data,p):return sorted([r for r in data['logs'] if r['project']==p['id']],key=lambda r:(r['date'],r['id']))
def actual(data,p):
    rows=records(data,p);volume=sum(r['volume_m3'] for r in rows);h=sum(r['heures_actives'] for r in rows)
    return dict(count=len(rows),volume=volume,hours=h,idle=sum(r['heures_arret'] for r in rows),fuel=sum(r['carburant_l'] for r in rows),
                cost=sum(r['depense_fcfa'] for r in rows),rate=volume/h if h else None,progress=100*volume/geometry(p['geometry'])['volume'])
def fingerprint(p):return hashlib.sha256(json.dumps(p,sort_keys=True).encode()).hexdigest()
def wilson(success,n):
    phat=success/n;z=1.95996398454;den=1+z*z/n
    center=(phat+z*z/(2*n))/den;delta=z*math.sqrt(phat*(1-phat)/n+z*z/(4*n*n))/den
    return [max(0,center-delta),min(1,center+delta)]


@st.cache_data(ttl=1800,max_entries=8,show_spinner=False)
def simulate(project_json):
    p=validate_project(json.loads(project_json));s=p['simulation'];m=next(m for m in p['machines'] if m['id']==p['selected']);g=p['geometry'];v=geometry(g)
    n=s['count'];rng=np.random.default_rng(s['seed'])
    if s['distribution']=='Fixe' or m['low']==m['high']:q=np.full(n,m['rate'])
    elif s['distribution']=='Uniforme':q=rng.uniform(m['low'],m['high'],n)
    else:q=rng.triangular(m['low'],m['rate'],m['high'],n)
    stopped=rng.random(n)<s['stop_probability'];draw=rng.uniform(s['stop_min'],s['stop_max'],n);idle=stopped*draw
    active=v['volume']/q;elapsed=active+idle;days=elapsed/g['hours_day'];fuel=active*m['fuel']
    cost=active*m['hourly']+idle*s['standby']*m['hourly']+g['setup']+v['loose']*g['haul']
    if not m['included']:cost+=fuel*m['fuel_price']
    over_time=days>p['deadline'];over_cost=cost>p['budget']
    sensitivity=[]
    for label,x in [('Rendement',q),('Durée d’arrêt',idle)]:
        for target,y in [('Durée',days),('Coût',cost)]:
            rho=float(stats.spearmanr(x,y).statistic) if np.ptp(x)>1e-12 and np.ptp(y)>1e-12 else None
            sensitivity.append(dict(facteur=label,indicateur=target,rho=rho))
    return dict(rate=q,active=active,idle=idle,days=days,fuel=fuel,cost=cost,
                p_time=float(over_time.mean()),p_cost=float(over_cost.mean()),p_either=float((over_time|over_cost).mean()),
                ci_time=wilson(int(over_time.sum()),n),ci_cost=wilson(int(over_cost.sum()),n),
                days_q=np.quantile(days,[.1,.5,.9]).tolist(),cost_q=np.quantile(cost,[.1,.5,.9]).tolist(),sensitivity=sensitivity)


def demo_store():
    p=new_project('Villa Horizon · Fondations','Yaoundé · exemple fictif');p['id']='TP-0001';p['start']=(date.today()-timedelta(days=3)).isoformat()
    p['soil']='Hétérogène / remblai';p['layers']=[dict(name='Terre de surface',thickness=.3),dict(name='Couche sableuse observée',thickness=1.2),dict(name='Couche non caractérisée',thickness=1.)]
    p['water']=2.;p['note']='Jeu fictif pour apprendre. Les couches, tarifs et rendements ne proviennent pas d’une étude locale.'
    q=new_project('Atelier Delta · Fouilles','Douala · exemple fictif');q['id']='TP-0002';q['geometry'].update(length=6.,width=3.,depth=1.2,count=2);q['budget']=400000.;q['deadline']=3.
    log=dict(id='DEMO-01',project=p['id'],date=date.today().isoformat(),engine='Pelle compacte',volume_m3=10.,heures_actives=2.,heures_arret=.5,carburant_l=8.,depense_fcfa=45000.,note='Mesures entièrement fictives ; 10 m³ réalisés.')
    return validate_store(dict(schema=1,projects=[p,q],logs=[log]))
def data_dir():
    default=Path.home()/('Documents/Codex/data/terrapilot' if os.name=='nt' else '.local/share/terrapilot')
    root=Path(os.environ.get('TERRAPILOT_DATA_DIR',str(default))).expanduser().resolve();app=Path(__file__).resolve().parent
    if root==app or app in root.parents:raise ValueError('Le registre doit être placé hors du dossier livré.')
    root.mkdir(parents=True,exist_ok=True);return root
def connect():
    conn=sqlite3.connect(data_dir()/'terrapilot.sqlite3',timeout=10)
    conn.execute('CREATE TABLE IF NOT EXISTS workspace(id INTEGER PRIMARY KEY CHECK(id=1),revision INTEGER NOT NULL,payload TEXT NOT NULL)')
    conn.execute('INSERT OR IGNORE INTO workspace VALUES(1,0,?)',(json.dumps(empty()),));conn.commit();return conn
def load_db():
    with closing(connect()) as conn:revision,payload=conn.execute('SELECT revision,payload FROM workspace WHERE id=1').fetchone()
    return validate_store(json.loads(payload)),revision
def save_db(data,revision):
    data=validate_store(data)
    with closing(connect()) as conn,conn:
        cur=conn.execute('UPDATE workspace SET revision=revision+1,payload=? WHERE id=1 AND revision=?',(json.dumps(data,ensure_ascii=False),revision))
        if cur.rowcount!=1:raise ValueError('Les données ont changé dans une autre session. Actualise puis recommence.')
def merge_backup(data,payload,allow_fictitious=False):
    if len(payload)>LIMIT:raise ValueError('Sauvegarde limitée à 15 Mo.')
    try:b=json.loads(payload)
    except Exception:raise ValueError('Fichier JSON illisible.') from None
    if not isinstance(b,dict) or b.get('application')!=APP:raise ValueError('Sauvegarde TerraPilot attendue.')
    if type(b.get('fictitious')) is not bool:raise ValueError('Origine de la sauvegarde manquante.')
    if b['fictitious'] and not allow_fictitious:raise ValueError('Restaurer les exemples dans l’espace fictif pour éviter de les mélanger aux mesures.')
    incoming=validate_store(b.get('data'));out=copy.deepcopy(data);mapping={}
    for p in incoming['projects']:
        old=p['id'];p['id']=code(out['projects']);mapping[old]=p['id'];p['name']=(p['name'][:91]+' (copie)');out['projects'].append(p)
    for r in incoming['logs']:r['project']=mapping[r['project']];r['id']=uid();out['logs'].append(r)
    return validate_store(out)
def log_signature(r):return tuple(str(r.get(k,'')) for k in LOG_COLUMNS)
def import_logs(data,p,raw):
    if len(raw)>2*1024*1024:raise ValueError('CSV limité à 2 Mo.')
    try:frame=pd.read_csv(io.BytesIO(raw),sep=None,engine='python',dtype=str,keep_default_na=False,encoding='utf-8-sig')
    except Exception:raise ValueError('CSV UTF-8 attendu, séparé par des virgules ou des points-virgules.') from None
    if set(frame.columns)!=set(LOG_COLUMNS) or not 1<=len(frame)<=1000:raise ValueError('Utiliser les 7 colonnes du modèle CSV, de 1 à 1 000 lignes.')
    result=[];seen={log_signature(r) for r in records(data,p)}
    for i,r in enumerate(frame.to_dict('records'),2):
        try:row=validate_log(dict(r,id=uid(),engine='Import CSV'),p)
        except ValueError as e:raise ValueError(f'Ligne {i} : {e}') from None
        key=log_signature(row)
        if key in seen:raise ValueError(f'Ligne {i} : relevé identique déjà présent. Aucun relevé importé.')
        seen.add(key);result.append(row)
    return result
def csv_bytes(frame):
    frame=frame.copy()
    for col in frame.select_dtypes(include=['object','str']).columns:
        frame[col]=frame[col].map(lambda v:"'"+v if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')) else v)
    return frame.to_csv(index=False,sep=';',decimal=',').encode('utf-8-sig')
def comparison(p):
    rows=[]
    for m in p['machines']:
        e=estimate(p,m);rows.append({'ID':m['id'],'Moyen':m['name'],'Rendement (m³/h)':m['rate'],'Durée active (h)':e['hours'],
                                   'Jours ouvrés équivalents':e['days'],'Carburant (L)':e['fuel'],'Coût total (FCFA)':e['total']})
    return pd.DataFrame(rows)


def fig_layout(fig,title='',height=340):
    fig.update_layout(title=dict(text=title,font=dict(size=16,color=GREEN)),height=height,paper_bgcolor='rgba(0,0,0,0)',plot_bgcolor='#FFFCF7',
                      font=dict(family='Segoe UI, Arial',color=GREEN,size=12),margin=dict(l=45,r=25,t=50 if title else 20,b=55),
                      legend=dict(orientation='h',y=-.2),colorway=[COPPER,GREEN,BLUE,'#CA9E6A'],hovermode='closest')
    fig.update_xaxes(gridcolor='#E6DECE',zeroline=False);fig.update_yaxes(gridcolor='#E6DECE',zeroline=False)
    return fig
def section_plot(p):
    g=p['geometry'];e=geometry(g);h=e['depth'];half=e['top_w']/2;span=max(half*1.5,2.);floor=max(h+.5,min(sum(x['thickness'] for x in p['layers']),h*1.7));fig=go.Figure();z=0
    palette=['#D6B98F','#BE9164','#9A7554','#7D6650']
    for i,layer in enumerate(p['layers']):
        bottom=min(z+layer['thickness'],floor)
        if z>=floor:break
        fig.add_shape(type='rect',x0=-span,x1=span,y0=-bottom,y1=-z,fillcolor=palette[i%4],line_width=0,layer='below')
        fig.add_trace(go.Scatter(x=[span*.82],y=[-(z+bottom)/2],mode='markers',marker=dict(size=5,color=GREEN),name=layer['name'],
                                 hovertemplate=esc(layer['name'])+'<br>Épaisseur saisie : '+fmt(layer['thickness'],2)+' m<extra></extra>',showlegend=False));z=bottom
    x=[-half,-e['bottom_w']/2,e['bottom_w']/2,half,-half];y=[0,-h,-h,0,0]
    fig.add_trace(go.Scatter(x=x,y=y,fill='toself',fillcolor='#FFF9EA',line=dict(color=COPPER,width=3),name='Fouille projetée',hovertemplate='Largeur : %{x:.2f} m<br>Cote : %{y:.2f} m<extra>Coupe géométrique</extra>'))
    if g['mode'].startswith('Semelle'):
        fig.add_shape(type='rect',x0=-g['width']/2,x1=g['width']/2,y0=-h,y1=-h+min(.2,h*.3),fillcolor='#6A9C9C',line_width=1,line_color=GREEN)
    if p['water'] is not None and p['water']<=floor:
        fig.add_trace(go.Scatter(x=[-span,span],y=[-p['water']]*2,line=dict(color=BLUE,dash='dash',width=2),name='Eau observée',hovertemplate='Profondeur eau : '+fmt(p['water'],2)+' m<extra></extra>'))
    fig.update_xaxes(title='Largeur de coupe (m)',range=[-span,span]);fig.update_yaxes(title='Cote relative (m)',range=[-floor,.45])
    fig.add_annotation(x=0,y=-h/2,text=f'H = {fmt(h,2)} m',showarrow=False,font=dict(color=GREEN,size=14))
    return fig_layout(fig,height=365)
def plan_plot(p):
    e=geometry(p['geometry']);fig=go.Figure()
    for name,l,w,color in [('Emprise en surface',e['top_l'],e['top_w'],COPPER),('Fond de fouille',e['bottom_l'],e['bottom_w'],GREEN)]:
        fig.add_trace(go.Scatter(x=[-l/2,l/2,l/2,-l/2,-l/2],y=[-w/2,-w/2,w/2,w/2,-w/2],mode='lines',line=dict(color=color,width=3),name=name))
    fig.update_xaxes(title='Longueur (m)');fig.update_yaxes(title='Largeur (m)',scaleanchor='x',scaleratio=1)
    return fig_layout(fig,height=330)


class PitDrawing(Flowable):
    def __init__(self,p):super().__init__();self.p=p;self.width=16*cm;self.height=4.8*cm
    def draw(self):
        c=self.canv;e=geometry(self.p['geometry']);w=self.width;h=self.height;scale=min((w-50)/e['top_w'],(h-35)/e['depth']);left=w/2
        c.setFillColor(colors.HexColor('#D6B98F'));c.rect(0,8,w,h-20,fill=1,stroke=0)
        path=c.beginPath();path.moveTo(left-e['top_w']*scale/2,h-12);path.lineTo(left-e['bottom_w']*scale/2,h-12-e['depth']*scale);path.lineTo(left+e['bottom_w']*scale/2,h-12-e['depth']*scale);path.lineTo(left+e['top_w']*scale/2,h-12);path.close()
        c.setFillColor(colors.HexColor('#FFFCF7'));c.setStrokeColor(colors.HexColor(COPPER));c.drawPath(path,fill=1,stroke=1)
        c.setFont('Terra',10);c.setFillColor(colors.HexColor(GREEN));c.drawCentredString(w/2,h/2,'Profondeur '+fmt(e['depth'],2)+' m')
        c.drawString(4,0,'Coupe géométrique illustrative : aucune stabilité des parois évaluée.')
def report_pdf(data,p,sim=None,fictitious=False):
    import reportlab
    folder=Path(reportlab.__file__).parent/'fonts'
    for name,file in [('Terra','Vera.ttf'),('TerraBold','VeraBd.ttf')]:
        if name not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont(name,str(folder/file)))
    pdfmetrics.registerFontFamily('Terra',normal='Terra',bold='TerraBold',italic='Terra',boldItalic='TerraBold')
    styles=getSampleStyleSheet()
    for s in styles.byName.values():s.fontName='TerraBold' if s.name in ('Title','Heading2') else 'Terra';s.textColor=colors.HexColor(GREEN)
    styles['BodyText'].leading=14
    styles['Heading2'].keepWithNext=1
    def P(text):return Paragraph(esc(text),styles['BodyText'])
    def T(rows,widths):
        t=Table([[P(x) for x in row] for row in rows],colWidths=widths,repeatRows=1,hAlign='LEFT')
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E1E9DE')),('VALIGN',(0,0),(-1,-1),'TOP'),
                               ('BOTTOMPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),8),('LINEBELOW',(0,0),(-1,-1),.3,colors.HexColor('#D8D0C0'))]))
        return t
    e=estimate(p);a=actual(data,p);body=[Paragraph(APP+' · Rapport de fouilles',styles['Title']),P(p['id']+' · '+p['name']),P(p['city']+' · '+date.today().isoformat()),P('EXEMPLE FICTIF' if fictitious else 'ESTIMATION ET RELEVÉS UTILISATEUR'),Spacer(1,.3*cm),
        P('Estimation pédagogique. Les tarifs, rendements, couches et pentes sont des données saisies. Ce rapport ne fixe ni une profondeur de fondation ni une stabilité de fouille.'),
        Paragraph('Géométrie et moyen retenu',styles['Heading2']),PitDrawing(p),Spacer(1,.3*cm)]
    rows=[['Indicateur','Valeur'],['Sol décrit',p['soil']],['Volume en place',fmt(e['volume'],3)+' m³'],['Volume foisonné',fmt(e['loose'],3)+' m³'],
          ['Moyen retenu',e['machine']],['Temps actif théorique',fmt(e['hours'],2)+' h'],['Jours ouvrés équivalents',fmt(e['days'],3)],['Consommation théorique',fmt(e['fuel'],2)+' L'],
          ['Coût théorique total',money(e['total'])],['Budget saisi',money(p['budget'])],['Échéance saisie',fmt(p['deadline'],2)+' jours ouvrés']]
    body+=[T(rows,[8*cm,8*cm]),Paragraph('Suivi mesuré / déclaré',styles['Heading2']),P(f"{a['count']} relevé(s) ; {fmt(a['volume'],2)} m³ réalisés ; {fmt(a['hours'],2)} h actives ; {money(a['cost'])} dépensés."),
           P('Ces relevés sont distincts des tirages simulés. Observations : '+(p['note'] or 'Aucune.'))]
    if sim is not None:
        body+=[Paragraph('Simulations conditionnelles aux hypothèses',styles['Heading2']),P(f"{p['simulation']['count']} tirages ; loi {p['simulation']['distribution']} ; graine {p['simulation']['seed']}. Un arrêt éventuel est tiré pour l’ensemble du chantier, pas pour chaque jour."),
               T([['Résultat','Valeur'],['P(durée > échéance)',fmt(100*sim['p_time'],2)+' %'],['P(coût > budget)',fmt(100*sim['p_cost'],2)+' %'],
                  ['Durée P50 / P90',fmt(sim['days_q'][1],3)+' / '+fmt(sim['days_q'][2],3)+' jours ouvrés'],['Coût P50 / P90',money(sim['cost_q'][1])+' / '+money(sim['cost_q'][2])]], [8*cm,8*cm])]
    body+=[Paragraph('Conventions',styles['Heading2']),P('Rendement en m³ de terrain en place par heure active. Les jours sont des équivalents ouvrés ; aucun calendrier météo réel n’est utilisé. Le foisonnement ne multiplie pas le temps d’excavation. Le carburant inclus n’est pas facturé deux fois. Les temps d’arrêt n’engendrent que la fraction du tarif horaire saisie pour l’attente.')]
    b=io.BytesIO()
    def footer(c,doc):c.setFont('Terra',8);c.drawString(1.6*cm,.8*cm,APP+' · Estimation pédagogique');c.drawRightString(19.3*cm,.8*cm,str(doc.page))
    SimpleDocTemplate(b,leftMargin=1.6*cm,rightMargin=1.6*cm,topMargin=1.5*cm,bottomMargin=1.5*cm,title=APP+' - Rapport').build(body,onFirstPage=footer,onLaterPages=footer)
    return b.getvalue()
def report_excel(data,p,sim=None,fictitious=False):
    e=estimate(p);wb=Workbook();s=wb.active;s.title='Projet'
    for row in [['TerraPilot','Valeur'],['Chantier',p['name']],['Volume en place (m³)',e['volume']],['Volume foisonné (m³)',e['loose']],['Temps actif (h)',e['hours']],['Coût théorique (FCFA)',e['total']],['Hypothèses',json.dumps(p['geometry'],ensure_ascii=False)]]:s.append(row)
    s.append(['Origine','EXEMPLE FICTIF' if fictitious else 'Données utilisateur'])
    s=wb.create_sheet('Engins');df=comparison(p);s.append(list(df.columns))
    for row in df.itertuples(index=False,name=None):s.append(list(row))
    s=wb.create_sheet('Journal');s.append(LOG_COLUMNS+['engine'])
    for r in records(data,p):s.append([r[c] for c in LOG_COLUMNS+['engine']])
    if sim is not None:
        s=wb.create_sheet('Simulation');s.append(['Indicateur','Valeur']);s.append(['Probabilité dépassement durée',sim['p_time']]);s.append(['Probabilité dépassement coût',sim['p_cost']]);s.append(['Durée P90 (jours ouvrés)',sim['days_q'][2]]);s.append(['Coût P90 (FCFA)',sim['cost_q'][2]])
    for s in wb:
        s.freeze_panes='A2'
        for cell in s[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='183F35')
        for column in s.columns:s.column_dimensions[column[0].column_letter].width=28
        for row in s:
            for cell in row:
                if isinstance(cell.value,str) and cell.value.lstrip().startswith(('=','+','-','@')):cell.data_type='s'
    b=io.BytesIO();wb.save(b);return b.getvalue()


CSS='''<style>
.stApp,[data-testid="stAppViewContainer"]{background:#F4F0E7;color:#183F35}
[data-testid="stHeader"]{background:#F4F0E7EC}.block-container{max-width:1450px;padding:3.6rem 2.2rem 2.5rem}
h1,h2,h3,h4{color:#183F35!important;letter-spacing:-.04em}h2{font-size:1.7rem!important}
[data-testid="stMain"] [data-testid="stMarkdownContainer"], [data-testid="stMain"] [data-testid="stWidgetLabel"] p{color:#25473D}
[data-testid="stMain"] [data-testid="stCaptionContainer"] p{color:#61756B!important}
[data-testid="stMain"] input,[data-testid="stMain"] textarea{background:#FFFCF7!important;color:#193E35!important}
[data-testid="stSidebar"]{background:#183F35;border-right:0}[data-testid="stSidebar"] p,[data-testid="stSidebar"] label{color:#E4EADD!important}
[data-testid="stSidebar"] button[kind="secondary"]{background:transparent;border:1px solid #3B5D4F;color:#F4F0E7}
[data-testid="stSidebar"] button[kind="primary"]{background:#E5BA88;border-color:#E5BA88;color:#183F35}
[data-testid="stSidebar"] button p{color:inherit!important}
[data-testid="stSidebar"] [data-testid="stWidgetLabel"] p{color:#C1D0BC!important}
[data-testid="stForm"]{background:#FFFCF7;border:1px solid #DED7C7;border-radius:17px}
[data-testid="stExpander"]{border:1px solid #DAD3C2;border-radius:13px;background:#FAF7EF}
.stButton button,.stFormSubmitButton button,.stDownloadButton button,.stLinkButton a{min-height:43px;border-radius:10px;font-weight:650}
[data-testid="stMain"] button[kind="primary"]{background:#B65C38;border-color:#B65C38;color:#fff}
[data-testid="stMain"] button[kind="primary"] [data-testid="stMarkdownContainer"]{color:#fff!important}
[data-testid="stMain"] button[kind="primaryFormSubmit"]{background:#B65C38;border-color:#B65C38;color:#fff}
[data-testid="stMain"] button[kind="primary"] p,[data-testid="stMain"] button[kind="primaryFormSubmit"] p{color:#fff!important}
[data-testid="stMain"] button[kind="secondary"]{background:#FFFCF7;color:#183F35;border-color:#D5CBB9}
button p{color:inherit!important}[data-testid="stDataFrame"]{border-radius:12px;border:1px solid #DAD3C2}
.tp-logo{display:flex;align-items:center;gap:11px;margin:12px 0 7px;color:#FAF4E7;font-size:30px;font-weight:850;letter-spacing:-1px}
.tp-symbol{display:inline-grid;place-items:center;background:#E5BA88;color:#183F35;border-radius:12px;width:39px;height:39px;font-size:24px;font-weight:900}
.tp-side-caption{color:#B8C9B2;font-size:11px;letter-spacing:1.5px;line-height:1.8;margin-bottom:26px}
.tp-side-note{border-top:1px solid #41604F;padding-top:15px;margin-top:26px;font-size:12px;color:#BECEB8;line-height:1.8}
.tp-hero{display:grid;grid-template-columns:1.3fr 1fr;gap:18px;padding:28px 31px;border-radius:22px;background:#183F35;color:#FAF5E9;overflow:hidden;margin-bottom:23px;position:relative}
.tp-hero:after{content:'';position:absolute;width:360px;height:360px;border:1px solid #41614E;border-radius:50%;right:-150px;top:-130px;pointer-events:none}
.tp-eyebrow{font:700 10px Arial;letter-spacing:2px;color:#D6B58A;margin-bottom:13px}
.tp-hero h1{font:800 clamp(27px,3.4vw,43px)/1.3 'Segoe UI',Arial;color:#FFFAEE!important;margin:0 0 13px;letter-spacing:-1.4px}
.tp-hero p{font-size:13px;line-height:1.8;color:#CFDCC7;margin:0;max-width:480px}
.tp-hero img{width:100%;height:190px;align-self:center;position:relative;z-index:1}
.tp-topline{display:flex;justify-content:space-between;align-items:center;gap:12px;margin-bottom:17px;color:#637568;font-size:12px;flex-wrap:wrap}
.tp-pill{font-size:11px;font-weight:700;background:#E6ECDC;border:1px solid #C4D0B8;color:#2F573E;border-radius:30px;padding:6px 12px}
.tp-pill.demo{background:#F4E3CB;color:#815137;border-color:#DAB88E}
.tp-kpis{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:13px;margin:15px 0 24px}
.tp-kpi{border-radius:16px;background:#FFFCF7;border:1px solid #E1D8C7;padding:20px 18px;min-width:0;border-top:3px solid #B65C38}
.tp-kpi.highlight{background:#E4EBD9;border-color:#CBD5BC;border-top-color:#4A7049}
.tp-kpi small{display:block;font-size:11px;color:#647767;line-height:1.5}.tp-kpi strong{display:block;font-size:clamp(22px,2.2vw,32px);letter-spacing:-1px;color:#183F35;margin:9px 0 4px;overflow-wrap:anywhere;line-height:1.3}.tp-kpi span{font-size:11px;color:#6B7C6F;line-height:1.5}
.tp-project{background:#FFFCF7;border:1px solid #DDD3C0;border-radius:16px;padding:20px;position:relative;min-height:185px;margin-bottom:8px}
.tp-project .id{font:700 10px Arial;color:#9A6446;letter-spacing:1px}.tp-project h3{font-size:20px;margin:10px 0 6px;padding:0}.tp-project p{font-size:12px;color:#687C6E;line-height:1.7;margin:0}
.tp-project .progress{display:flex;align-items:center;gap:13px;margin-top:18px}
.tp-ring{width:56px;height:56px;border-radius:50%;background:conic-gradient(#B65C38 var(--p),#ECE5D6 0);display:grid;place-items:center;flex-shrink:0}
.tp-ring b{width:44px;height:44px;border-radius:50%;background:#FFFCF7;display:grid;place-items:center;font-size:12px;color:#183F35}
.tp-panel-title{display:flex;justify-content:space-between;align-items:center;gap:12px;margin:8px 0 9px}.tp-panel-title h3{font-size:21px;margin:0;padding:0}.tp-panel-title span{font-size:11px;color:#78856F}
.tp-note{border-left:3px solid #B65C38;background:#EBE5D7;color:#4A5D4B;padding:13px 16px;font-size:12px;line-height:1.8;border-radius:0 10px 10px 0;margin:12px 0}
.st-key-cut,.st-key-distribution,.st-key-journal_chart{background:#FFFCF7;border:1px solid #DED6C6;border-radius:18px;padding:15px}
.tp-footer{font-size:10px;color:#7E8776;margin-top:32px;padding-top:15px;border-top:1px solid #DDD4C2;line-height:1.8}
@media(max-width:1000px){.tp-kpis{grid-template-columns:repeat(2,minmax(0,1fr))}.tp-hero{padding:23px;grid-template-columns:1.1fr 1fr}.tp-hero h1{font-size:33px}}
@media(max-width:640px){.block-container{padding:3.5rem .85rem 2rem}.tp-hero{grid-template-columns:1fr;padding:22px 20px;gap:7px;border-radius:16px}.tp-hero h1{font-size:30px}.tp-hero img{height:135px}.tp-kpis{gap:9px}.tp-kpi{padding:15px 12px}.tp-kpi strong{font-size:23px}.tp-kpi span{font-size:10px}.tp-project{padding:17px}.tp-panel-title h3{font-size:18px}}
@media(prefers-reduced-motion:reduce){*{animation:none!important;transition:none!important}}
</style>'''

HERO_SVG='''<svg viewBox="0 0 450 210" xmlns="http://www.w3.org/2000/svg" role="img" aria-label="Illustration d’une pelle et d’une coupe de fouille">
<defs><pattern id="soil" width="18" height="12" patternUnits="userSpaceOnUse"><path d="M2 4h4m7 6h3" stroke="#987755" stroke-width="1"/></pattern></defs>
<path d="M15 102H430V193H15Z" fill="#BC966A"/><path d="M15 134H430V193H15Z" fill="#8F7657"/><path d="M15 164H430V193H15Z" fill="#645F48"/><path d="M15 104H430V193H15Z" fill="url(#soil)"/>
<path d="M190 102L220 164H333L370 102" fill="#183F35" stroke="#E5BA88" stroke-width="2"/><path d="M225 167H330" stroke="#B6D1B7" stroke-width="6"/>
<path d="M74 57h62l18 31H62Z" fill="#D99158"/><rect x="98" y="34" width="31" height="30" rx="4" fill="#E8BD87"/><path d="M103 39h20v16h-20Z" fill="#426D61"/>
<rect x="55" y="87" width="105" height="17" rx="8" fill="#253E32" stroke="#E9C897" stroke-width="2"/><path d="M75 96H142" stroke="#BD9971" stroke-width="3" stroke-dasharray="6 6"/>
<path d="M141 66L193 27L248 83L277 112" fill="none" stroke="#DFA76C" stroke-width="13" stroke-linecap="round"/><circle cx="192" cy="28" r="8" fill="#E8C593"/><path d="M266 102L292 113L284 135L261 123Z" fill="#D28550"/>
<path d="M396 109V161M390 109h12M390 161h12" stroke="#D7DDC1" stroke-width="1.3"/><text x="407" y="140" fill="#EAE8CF" font-family="Arial" font-size="13">H</text>
<path d="M16 101h157m205 0h54" stroke="#CBDBBB" stroke-width="3"/><path d="M27 62h48m-24-12h23" stroke="#4A6B54" stroke-width="2"/>
</svg>'''


def current():return st.session_state['tp_demo'] if st.session_state['tp_space']=='Exemple fictif' else st.session_state['tp_real']
def active(data):return next((p for p in data['projects'] if p['id']==st.session_state.get('tp_active')),None)
def nav(page):st.session_state['tp_page']=page
def choose(pid):st.session_state['tp_active']=pid;nav('Vue d’ensemble')
def load_demo():
    st.session_state['tp_demo']=demo_store();st.session_state['tp_space']='Exemple fictif';st.session_state['_tp_space']='Exemple fictif';st.session_state['tp_active']='TP-0001';nav('Vue d’ensemble')
def switch_space():st.session_state['tp_space']=st.session_state['_tp_space'];st.session_state.pop('tp_active',None)
def commit(data,message):
    try:
        data=validate_store(data)
        if st.session_state['tp_space']=='Exemple fictif':st.session_state['tp_demo']=data
        else:save_db(data,st.session_state['tp_revision'])
        st.session_state['tp_flash']=message;st.rerun()
    except (ValueError,OSError,sqlite3.Error) as e:st.error(str(e))
def save_project(data,p,message):
    out=copy.deepcopy(data);out['projects']=[p if r['id']==p['id'] else r for r in out['projects']];commit(out,message)
def require(p):
    if p:return True
    st.info('Crée un chantier ou charge l’exemple fictif pour utiliser cet outil.')
    if st.button('Créer mon premier chantier',type='primary'):nav('Chantiers');st.rerun()
    return False
def cached_sim(p):
    stored=st.session_state.get('tp_result',{})
    return stored.get('result') if stored.get('hash')==fingerprint(p) else None
def kpis(items):
    st.html('<div class="tp-kpis">'+''.join('<div class="tp-kpi'+(' highlight' if i==0 else '')+'"><small>'+esc(label)+'</small><strong>'+esc(value)+'</strong><span>'+esc(sub)+'</span></div>' for i,(label,value,sub) in enumerate(items))+'</div>')
def panel(title,sub=''):st.html('<div class="tp-panel-title"><h3>'+esc(title)+'</h3><span>'+esc(sub)+'</span></div>')
def project_cards(data):
    search=st.text_input('Rechercher un chantier',placeholder='Nom, localisation…',key='tp_search')
    rows=[p for p in data['projects'] if search.casefold() in (p['name']+' '+p['city']).casefold()]
    if not rows:st.caption('Aucun chantier ne correspond à cette recherche.')
    for i in range(0,len(rows),3):
        for col,p in zip(st.columns(3),rows[i:i+3]):
            a=actual(data,p);e=estimate(p);progress=min(a['progress'],100)
            with col:
                st.html('<article class="tp-project"><div class="id">'+esc(p['id'])+'</div><h3>'+esc(p['name'])+'</h3><p>'+esc(p['city'])+'</p><div class="progress"><div class="tp-ring" style="--p:'+str(progress)+'%"><b>'+fmt(a['progress'])+'%</b></div><p>'+fmt(a['volume'],1)+' / '+fmt(e['volume'],1)+' m³<br>'+money(a['cost'])+' déclarés</p></div></article>')
                st.button('Ouvrir '+p['id'],key='open_'+p['id'],on_click=choose,args=(p['id'],),width='stretch',icon=':material/arrow_outward:')


def dashboard(data,p):
    illustration='<img alt="Pelle et coupe de fouille illustratives" src="data:image/svg+xml;base64,'+base64.b64encode(HERO_SVG.encode()).decode()+'">'
    st.html('<section class="tp-hero"><div><div class="tp-eyebrow">FOUILLES · ENGINS · INCERTITUDES</div><h1>Préparer les fouilles.<br>Piloter le terrain.</h1><p>Du volume à excaver au dernier relevé : une lecture claire des moyens, des délais et des dépenses.</p></div>'+illustration+'</section>')
    if not p:
        kpis([('Chantiers','0','Ton espace est prêt'),('Volumes','—','Des dimensions aux m³'),('Prévisions','—','Des hypothèses aux scénarios'),('Suivi','—','Des relevés aux décisions')])
        c,d=st.columns(2)
        with c:
            st.subheader('Commencer sur le terrain')
            st.write('Crée une fiche de chantier, saisis tes dimensions, puis compare les moyens de travail.')
            st.button('Créer un chantier',type='primary',on_click=nav,args=('Chantiers',),icon=':material/add:')
        with d:
            st.subheader('Découvrir avec un exemple')
            st.write('Deux chantiers fictifs, des engins comparables et un relevé de terrain pour essayer chaque outil.')
            st.button('Explorer l’exemple fictif',on_click=load_demo,icon=':material/science:')
        return
    e=estimate(p);a=actual(data,p);sim=cached_sim(p)
    kpis([('À excaver · en place',fmt(e['volume'],1)+' m³',fmt(e['loose'],1)+' m³ foisonnés à transporter'),
          ('Temps actif théorique',fmt(e['hours'],2)+' h',fmt(e['days'],2)+' jours ouvrés équivalents'),
          ('Estimation de base',money(e['total']),e['machine']+' · hors arrêt aléatoire'),
          ('Avancement déclaré',fmt(a['progress'],1)+' %',fmt(a['volume'],1)+' m³ réellement saisis')])
    left,right=st.columns([1.3,1])
    with left:
        with st.container(key='cut'):
            panel('Lire le terrain','Coupe transversale')
            st.plotly_chart(section_plot(p),width='stretch',alt='Coupe géométrique de la fouille, couches observées et éventuelle eau observée',key='dash_cut')
            st.caption('Couches décrites par l’utilisateur ; dessin illustratif, sans calcul de stabilité.')
    with right:
        panel('Dépenser avec une vue d’ensemble')
        fig=go.Figure(go.Bar(x=['Budget saisi','Base théorique','Réel déclaré'],y=[p['budget'],e['total'],a['cost']],marker_color=['#779566',COPPER,GREEN],text=[money(p['budget']),money(e['total']),money(a['cost'])],textposition='auto'))
        fig.update_yaxes(title='FCFA');st.plotly_chart(fig_layout(fig,height=365),width='stretch',alt='Comparaison du budget, du coût théorique et des dépenses déclarées')
    with st.container(horizontal=True):
        st.button('Ajuster la fouille et les engins',on_click=nav,args=('Fouilles et engins',),icon=':material/tune:')
        st.button('Simuler les incertitudes',on_click=nav,args=('Simulations',),type='primary',icon=':material/play_arrow:')
        st.button('Ajouter un relevé',on_click=nav,args=('Suivi des travaux',),icon=':material/edit_note:')
    if sim is not None:st.caption('Dernière simulation de ces paramètres : '+fmt(sim['p_time']*100,1)+' % de dépassements de délai ; '+fmt(sim['p_cost']*100,1)+' % de dépassements de budget.')
    panel('Tes chantiers','Ouvrir une carte pour changer de chantier')
    project_cards(data)


def project_form(data,p=None):
    editing=p is not None;draft=copy.deepcopy(p or new_project());key='edit_'+draft['id'] if editing else 'new'
    with st.form('project_'+key,clear_on_submit=not editing):
        name=st.text_input('Nom du chantier',value=draft['name'] if editing else '',max_chars=100)
        city=st.text_input('Localisation',value=draft['city'] if editing else '',max_chars=100)
        c,d=st.columns(2)
        with c:
            start=st.date_input('Début du chantier',value=iso(draft['start']))
            soil=st.selectbox('Nature du sol observée',SOILS,index=SOILS.index(draft['soil']))
            budget=st.number_input('Budget de fouille (FCFA)',0.,1e12,float(draft['budget']),step=10000.)
        with d:
            deadline=st.number_input('Échéance (jours ouvrés équivalents)',.01,10000.,float(draft['deadline']),step=.25)
            known=st.checkbox('Profondeur d’eau observée connue',value=draft['water'] is not None)
            water=st.number_input('Profondeur d’eau observée (m)',0.,1000.,float(draft['water'] or 0.),step=.1)
        note=st.text_area('Observations du terrain',value=draft['note'],max_chars=2000,placeholder='Accès, présence d’eau, limites de l’observation…')
        layers=st.data_editor(pd.DataFrame(draft['layers']),key='layers_'+key,num_rows='dynamic',hide_index=True,width='stretch',
                              column_config={'name':st.column_config.TextColumn('Couche observée',required=True),'thickness':st.column_config.NumberColumn('Épaisseur (m)',min_value=.01,max_value=1000.,required=True)},
                              alt='Tableau modifiable des couches de sol décrites')
        photo=st.file_uploader('Photo du chantier (facultatif)',type=['jpg','jpeg','png','webp'],max_upload_size=5,key='photo_'+key)
        submitted=st.form_submit_button('Enregistrer la fiche' if editing else 'Créer le chantier',type='primary')
    if submitted:
        try:
            draft.update(name=name,city=city,start=start.isoformat(),soil=soil,budget=budget,deadline=deadline,water=water if known else None,note=note,layers=layers.to_dict('records'))
            if photo:draft['photo']=photo_bytes(photo.getvalue())
            if editing:save_project(data,draft,'Fiche mise à jour.')
            else:
                out=copy.deepcopy(data);draft['id']=code(out['projects']);out['projects'].append(draft);st.session_state['tp_active']=draft['id'];commit(out,'Chantier créé.')
        except ValueError as error:st.error(str(error))


def projects_page(data,p):
    st.header('Chantiers');st.caption('Une carte par chantier. Les dimensions, engins, simulations et relevés appartiennent à la fiche ouverte.')
    with st.expander('Créer un nouveau chantier',expanded=not data['projects'],key='ex_new_project'):project_form(data)
    if data['projects']:project_cards(data)
    if p:
        st.subheader('Fiche active · '+p['name'])
        if p['photo']:st.image(base64.b64decode(p['photo']),width='stretch',alt='Photo du chantier '+p['name'])
        with st.expander('Modifier la fiche active',key='ex_edit_project_'+p['id']):project_form(data,p)
        st.caption('La nature du sol est une description saisie : elle ne fournit pas automatiquement une portance ou une pente admissible.')


def geometry_page(data,p):
    st.header('Fouilles et engins')
    if not require(p):return
    st.caption('Les dimensions donnent le volume. Les rendements et les tarifs donnent le temps et le coût.')
    g=p['geometry'];key=p['id']
    with st.form('geom_'+key):
        mode=st.selectbox('Type de saisie',['Fouille rectangulaire','Semelle avec espace de travail'],index=['Fouille rectangulaire','Semelle avec espace de travail'].index(g['mode']))
        st.caption('En mode semelle, longueur et largeur décrivent la semelle ; l’espace de travail est ajouté de chaque côté. La profondeur reste celle de la fouille.')
        a,b,c=st.columns(3)
        with a:
            length=st.number_input('Longueur (m)',.01,1000.,float(g['length']),step=.1)
            count=st.number_input('Nombre de fouilles identiques',1,1000,int(g['count']))
            hours=st.number_input('Heures disponibles par jour',.1,24.,float(g['hours_day']),step=.5)
        with b:
            width=st.number_input('Largeur (m)',.01,1000.,float(g['width']),step=.1)
            clearance=st.number_input('Espace de travail par côté (m)',0.,10.,float(g['clearance']),step=.1)
            setup=st.number_input('Frais fixes / mobilisation (FCFA)',0.,1e10,float(g['setup']),step=1000.)
        with c:
            depth=st.number_input('Profondeur de fouille (m)',.01,30.,float(g['depth']),step=.1)
            slope=st.number_input('Talus géométrique (m horizontal / m vertical)',0.,3.,float(g['slope']),step=.1)
            bulking=st.number_input('Foisonnement saisi (%)',0.,100.,float(g['bulking']),step=1.)
        haul=st.number_input('Transport par m³ foisonné (FCFA/m³)',0.,1e7,float(g['haul']),step=500.)
        submitted=st.form_submit_button('Calculer et enregistrer la fouille',type='primary')
    if submitted:
        updated=copy.deepcopy(p);updated['geometry']=dict(mode=mode,length=length,width=width,depth=depth,count=count,clearance=clearance,slope=slope,bulking=bulking,hours_day=hours,setup=setup,haul=haul)
        save_project(data,updated,'Géométrie enregistrée. Relancer la simulation si ses hypothèses ont changé.')
    e=estimate(p);kpis([('Volume en place',fmt(e['volume'],3)+' m³','Volume réellement à excaver'),('Volume foisonné',fmt(e['loose'],3)+' m³','Volume théorique à transporter'),('Emprise en surface',fmt(e['surface'],2)+' m²','Somme des emprises des fouilles'),('Temps actif',fmt(e['hours'],2)+' h',e['machine'])])
    a,b=st.columns(2)
    with a:panel('Coupe de fouille');st.plotly_chart(section_plot(p),width='stretch',alt='Coupe de fouille avec profondeur et couches de sol')
    with b:panel('Vue en plan · une fouille');st.plotly_chart(plan_plot(p),width='stretch',alt='Emprise de la fouille en surface et au fond')
    st.caption('Le talus saisi décrit uniquement la géométrie. Ni la stabilité des parois ni la profondeur des fondations ne sont déterminées par cet outil.')
    panel('Comparer les moyens de travail','Hypothèses fictives initiales, modifiables')
    frame=comparison(p);st.dataframe(frame.drop(columns='ID'),hide_index=True,width='stretch',alt='Comparaison des durées, consommations et coûts des engins')
    fig=go.Figure(go.Bar(x=frame['Moyen'],y=frame['Coût total (FCFA)'],marker_color=[COPPER if i==p['selected'] else '#8A9D7B' for i in frame['ID']]))
    fig.update_yaxes(title='FCFA');st.plotly_chart(fig_layout(fig,height=300),width='stretch',alt='Coût théorique de chaque moyen de travail')
    with st.form('machine_select_'+key):
        selected=st.selectbox('Moyen retenu pour les prévisions',[m['id'] for m in p['machines']],index=[m['id'] for m in p['machines']].index(p['selected']),format_func=lambda x:next(m['name'] for m in p['machines'] if m['id']==x))
        select=st.form_submit_button('Retenir ce moyen',type='primary')
    if select:
        updated=copy.deepcopy(p);updated['selected']=selected;save_project(data,updated,'Moyen retenu enregistré.')
    with st.expander('Modifier ou ajouter un moyen de travail',key='ex_machine_'+key):
        machine_id=st.selectbox('Fiche du moyen de travail',[m['id'] for m in p['machines']]+['nouveau'],format_func=lambda x:'Ajouter un moyen' if x=='nouveau' else next(m['name'] for m in p['machines'] if m['id']==x),key='editmachine_'+key)
        m=next((r for r in p['machines'] if r['id']==machine_id),dict(id=uid(),name='',low=4.,rate=5.,high=6.,hourly=20000.,fuel=4.,fuel_price=1000.,included=True))
        with st.form('machine_'+key+'_'+machine_id,clear_on_submit=machine_id=='nouveau'):
            name=st.text_input('Nom du moyen de travail',value=m['name'],max_chars=80)
            a,b,c=st.columns(3)
            with a:low=st.number_input('Rendement minimum (m³/h)',.01,10000.,float(m['low']));hourly=st.number_input('Tarif horaire (FCFA/h)',0.,1e8,float(m['hourly']),step=1000.)
            with b:rate=st.number_input('Rendement habituel (m³/h)',.01,10000.,float(m['rate']));fuel=st.number_input('Consommation (L/h active)',0.,1000.,float(m['fuel']))
            with c:high=st.number_input('Rendement maximum (m³/h)',.01,10000.,float(m['high']));fuel_price=st.number_input('Prix du carburant (FCFA/L)',0.,10000.,float(m['fuel_price']),step=50.)
            included=st.checkbox('Carburant inclus dans le tarif horaire',value=m['included'])
            submit=st.form_submit_button('Enregistrer ce moyen')
        if submit:
            updated=copy.deepcopy(p);edited=dict(id=m['id'],name=name,low=low,rate=rate,high=high,hourly=hourly,fuel=fuel,fuel_price=fuel_price,included=included)
            if machine_id=='nouveau':updated['machines'].append(edited)
            else:updated['machines']=[edited if r['id']==machine_id else r for r in updated['machines']]
            save_project(data,updated,'Moyen de travail enregistré.')
    with st.expander('Comprendre le calcul'):
        st.latex(r'V=\frac{H}{6}(A_{fond}+4A_{milieu}+A_{surface})\times N')
        st.write('Les côtés varient linéairement avec la profondeur. Sans talus, V = L × l × H × N. L’espace de travail est ajouté seulement en mode semelle.')
        st.latex(r't_{actif}=\frac{V_{en\ place}}{Q}\quad ;\quad V_{foisonne}=V_{en\ place}(1+f/100)')
        st.write('Coût = temps actif × tarif horaire + carburant non inclus + transport du volume foisonné + frais fixes. Les heures d’arrêt sont ajoutées seulement dans les simulations ou les relevés.')


def simulations_page(data,p):
    st.header('Simulations des délais et des coûts')
    if not require(p):return
    e=estimate(p);s=p['simulation'];m=next(m for m in p['machines'] if m['id']==p['selected'])
    st.caption('Un tirage représente un chantier complet, avec un rendement global et au plus un arrêt agrégé. Aucune météo réelle ni probabilité d’effondrement n’est calculée.')
    st.write('Moyen retenu : **'+m['name']+'** · rendement minimum / habituel / maximum : **'+fmt(m['low'],2)+' / '+fmt(m['rate'],2)+' / '+fmt(m['high'],2)+' m³/h**.')
    with st.form('simulation_'+p['id']):
        a,b,c=st.columns(3)
        with a:
            distribution=st.selectbox('Loi du rendement',['Triangulaire','Uniforme','Fixe'],index=['Triangulaire','Uniforme','Fixe'].index(s['distribution']))
            count=st.selectbox('Nombre de scénarios',[1000,5000,10000,50000,100000],index=[1000,5000,10000,50000,100000].index(s['count']) if s['count'] in [1000,5000,10000,50000,100000] else 2)
        with b:
            probability=st.number_input('Probabilité d’un arrêt agrégé (%)',0.,100.,s['stop_probability']*100,step=5.)
            stop_min=st.number_input('Arrêt minimum si événement (h)',0.,500.,s['stop_min'],step=.5)
        with c:
            standby=st.number_input('Tarif d’attente (% du tarif horaire)',0.,100.,s['standby']*100,step=5.)
            stop_max=st.number_input('Arrêt maximum si événement (h)',0.,500.,s['stop_max'],step=.5)
        seed=st.number_input('Graine aléatoire reproductible',0,2**32-1,int(s['seed']))
        submitted=st.form_submit_button('Lancer les simulations',type='primary',icon=':material/play_arrow:')
    if submitted:
        updated=copy.deepcopy(p);updated['simulation']=dict(distribution=distribution,count=count,seed=seed,stop_probability=probability/100,stop_min=stop_min,stop_max=stop_max,standby=standby/100)
        try:
            updated=validate_project(updated)
            with st.spinner('Calcul des scénarios…'):result=simulate(json.dumps(updated,sort_keys=True))
            out=copy.deepcopy(data);out['projects']=[updated if x['id']==p['id'] else x for x in out['projects']]
            # Résultats de session : aucune ligne simulée n’entre dans le journal mesuré.
            st.session_state['tp_result']=dict(hash=fingerprint(updated),result=result)
            commit(out,'Simulation terminée. Les résultats dépendent des hypothèses saisies.')
        except ValueError as error:st.error(str(error))
    sim=cached_sim(p)
    if sim is None:
        st.info('Lance une simulation. Après une modification de la fiche, des dimensions ou des engins, les anciens résultats sont masqués.')
        return
    kpis([('Dépassement de délai',fmt(100*sim['p_time'],1)+' %','Échéance : '+fmt(p['deadline'],2)+' jours ouvrés'),
          ('Dépassement de budget',fmt(100*sim['p_cost'],1)+' %','Budget : '+money(p['budget'])),
          ('Durée P90',fmt(sim['days_q'][2],2)+' j','90 % des tirages sont en dessous ou à cette durée'),
          ('Coût P90',money(sim['cost_q'][2]),'90 % des tirages sont en dessous ou à ce coût')])
    a,b=st.columns(2)
    with a:
        fig=go.Figure(go.Histogram(x=sim['days'],nbinsx=45,marker_color=COPPER,name='Durées simulées'))
        fig.add_vline(x=p['deadline'],line_dash='dash',line_color=GREEN);fig.update_xaxes(title='Jours ouvrés équivalents');fig.update_yaxes(title='Nombre de scénarios')
        st.plotly_chart(fig_layout(fig,'Distribution des durées'),width='stretch',alt='Histogramme des durées simulées et échéance saisie')
    with b:
        ordered=np.sort(sim['cost']);indices=np.linspace(0,len(ordered)-1,min(800,len(ordered))).astype(int)
        fig=go.Figure(go.Scatter(x=ordered[indices],y=(indices+1)/len(ordered)*100,mode='lines',line=dict(color=GREEN,width=3),fill='tozeroy',fillcolor='#DCE5D2'))
        fig.add_vline(x=p['budget'],line_dash='dash',line_color=COPPER);fig.update_xaxes(title='Coût total (FCFA)');fig.update_yaxes(title='Scénarios à ce coût ou moins (%)')
        st.plotly_chart(fig_layout(fig,'Courbe cumulée des coûts'),width='stretch',alt='Distribution cumulée empirique du coût simulé et budget saisi')
    table=pd.DataFrame({'Quantile':['P10','P50 · médiane','P90'],'Durée (jours ouvrés)':sim['days_q'],'Coût (FCFA)':sim['cost_q']})
    st.dataframe(table,hide_index=True,width='stretch',alt='Quantiles 10, 50 et 90 pour le coût et la durée')
    panel('Quels facteurs font varier les résultats ?','Corrélation de rang de Spearman')
    sensitivity=pd.DataFrame(sim['sensitivity']);st.dataframe(sensitivity,hide_index=True,width='stretch',alt='Sensibilité descriptive entre rendement, arrêts, durée et coût')
    st.caption('ρ proche de −1 : le résultat diminue quand le facteur augmente ; proche de +1 : il augmente. Valeur vide si le facteur ou le résultat est constant. Cette corrélation n’est pas une part causale de variance.')
    st.write('Intervalle de Wilson à 95 % pour la fréquence de dépassement du délai : **'+fmt(100*sim['ci_time'][0],2)+' à '+fmt(100*sim['ci_time'][1],2)+' %**. Pour le budget : **'+fmt(100*sim['ci_cost'][0],2)+' à '+fmt(100*sim['ci_cost'][1],2)+' %**.')
    st.caption('Ces intervalles reflètent l’échantillonnage Monte-Carlo, pas l’incertitude sur la validité du modèle. Zéro dépassement simulé ne prouve pas un risque réel nul.')
    frame=pd.DataFrame({'rendement_m3_h':sim['rate'],'heures_actives':sim['active'],'heures_arret':sim['idle'],'jours_ouvres':sim['days'],'carburant_l':sim['fuel'],'cout_fcfa':sim['cost']})
    st.download_button('Télécharger les scénarios CSV',csv_bytes(frame),'terrapilot_scenarios.csv','text/csv',icon=':material/download:')
    with st.expander('Lire les hypothèses'):
        st.write('Triangulaire : minimum, valeur la plus probable et maximum. Uniforme : toutes les valeurs entre minimum et maximum ont la même densité. Fixe : rendement habituel sans variabilité de rendement.')
        st.write('Un arrêt survient avec la probabilité saisie ; sa durée conditionnelle est uniforme entre les deux bornes. Il est indépendant du rendement. Le carburant est consommé pendant le travail actif seulement. Le coût d’attente est une fraction du tarif horaire ; ni le carburant ni les frais fixes ne sont ajoutés une deuxième fois.')
        st.write('Les prévisions portent sur l’intégralité du volume planifié. Elles ne sont pas une prévision automatique du reste à faire après les relevés.')


def add_or_edit_log(data,p,row):
    try:
        row=validate_log(row,p);out=copy.deepcopy(data)
        if any(log_signature(r)==log_signature(row) and r['id']!=row['id'] for r in records(data,p)):raise ValueError('Un relevé identique existe déjà.')
        if any(r['id']==row['id'] for r in out['logs']):out['logs']=[row if r['id']==row['id'] else r for r in out['logs']]
        else:out['logs'].append(row)
        commit(out,'Relevé enregistré.')
    except ValueError as error:st.error(str(error))
def log_form(data,p,row=None):
    editing=row is not None;r=row or dict(date=date.today().isoformat(),volume_m3=0.,heures_actives=0.,heures_arret=0.,carburant_l=0.,depense_fcfa=0.,note='',engine=next(m['name'] for m in p['machines'] if m['id']==p['selected']))
    with st.form('log_'+p['id']+'_'+(row['id'] if editing else 'new'),clear_on_submit=not editing):
        a,b,c=st.columns(3)
        with a:d=st.date_input('Date du relevé',value=iso(r['date']));volume=st.number_input('Volume excavé ce relevé (m³ en place)',0.,1e7,float(r['volume_m3']),step=1.)
        with b:hours=st.number_input('Heures actives de ce relevé',0.,24.,float(r['heures_actives']),step=.5);idle=st.number_input('Heures d’arrêt de ce relevé',0.,24.,float(r['heures_arret']),step=.5)
        with c:fuel=st.number_input('Carburant consommé ce relevé (L)',0.,1e6,float(r['carburant_l']),step=1.);cost=st.number_input('Dépense totale de ce relevé (FCFA)',0.,1e12,float(r['depense_fcfa']),step=1000.)
        engine=st.text_input('Moyen réellement utilisé',value=r['engine'],max_chars=80)
        note=st.text_area('Note du relevé',value=r['note'],max_chars=1000)
        st.caption('Saisir les quantités de cette intervention, pas les cumuls. La dépense totale provient de tes relevés ; les litres ne sont pas refacturés automatiquement.')
        submitted=st.form_submit_button('Mettre à jour le relevé' if editing else 'Enregistrer le relevé',type='primary')
    if submitted:add_or_edit_log(data,p,dict(id=row['id'] if editing else uid(),date=d.isoformat(),volume_m3=volume,heures_actives=hours,heures_arret=idle,carburant_l=fuel,depense_fcfa=cost,note=note,engine=engine))
def tracking_page(data,p):
    st.header('Suivi des travaux')
    if not require(p):return
    st.caption('Les relevés sont mesurés ou déclarés. Aucun tirage Monte-Carlo n’est ajouté à ce journal.')
    a=actual(data,p);e=estimate(p)
    kpis([('Volume réalisé',fmt(a['volume'],2)+' m³',fmt(a['progress'],1)+' % du volume planifié'),('Dépenses déclarées',money(a['cost']),'Budget : '+money(p['budget'])),
          ('Heures actives',fmt(a['hours'],2)+' h',fmt(a['idle'],2)+' h d’arrêt déclarées'),('Rendement observé',fmt(a['rate'],2)+' m³/h' if a['rate'] is not None else '—',str(a['count'])+' relevé(s)')])
    if a['progress']>100:st.info('Le volume déclaré dépasse le volume planifié. Vérifie le métré et les relevés ; le dépassement est conservé.')
    with st.expander('Ajouter un relevé de terrain',expanded=True,key='ex_log_add_'+p['id']):log_form(data,p)
    rows=records(data,p)
    if rows:
        frame=pd.DataFrame(rows);daily=frame.groupby('date',as_index=False)[['volume_m3','depense_fcfa','heures_actives','heures_arret']].sum()
        daily['volume_cumule']=daily['volume_m3'].cumsum();daily['cout_cumule']=daily['depense_fcfa'].cumsum()
        a,b=st.columns(2)
        with a:
            fig=go.Figure(go.Scatter(x=daily['date'],y=daily['volume_cumule'],mode='lines+markers',line=dict(color=GREEN,width=3),fill='tozeroy',fillcolor='#DDE6D2'))
            fig.add_hline(y=e['volume'],line_dash='dash',line_color=COPPER);fig.update_xaxes(type='category',title='Date de relevé');fig.update_yaxes(title='m³ en place cumulés');st.plotly_chart(fig_layout(fig,'Avancement des fouilles'),width='stretch',alt='Volumes réellement déclarés cumulés et objectif planifié')
        with b:
            fig=go.Figure(go.Scatter(x=daily['date'],y=daily['cout_cumule'],mode='lines+markers',line=dict(color=COPPER,width=3)))
            fig.add_hline(y=p['budget'],line_dash='dash',line_color=GREEN);fig.update_xaxes(type='category',title='Date de relevé');fig.update_yaxes(title='FCFA cumulés');st.plotly_chart(fig_layout(fig,'Évolution financière réelle'),width='stretch',alt='Dépenses déclarées cumulées comparées au budget')
        panel('Journal des interventions');st.dataframe(frame[['date','engine']+LOG_COLUMNS[1:]],hide_index=True,width='stretch',alt='Journal des relevés réels saisis')
        with st.expander('Corriger ou supprimer un relevé',key='ex_log_edit_'+p['id']):
            selected=st.selectbox('Relevé à corriger',[r['id'] for r in rows],format_func=lambda i:next(r['date']+' · '+fmt(r['volume_m3'],2)+' m³ · '+r['engine'] for r in rows if r['id']==i),key='logselect_'+p['id'])
            row=next(r for r in rows if r['id']==selected);log_form(data,p,row)
            confirm=st.checkbox('Confirmer la suppression de ce relevé',key='confirm_'+selected)
            if st.button('Supprimer le relevé sélectionné',disabled=not confirm,key='delete_'+selected):
                out=copy.deepcopy(data);out['logs']=[r for r in out['logs'] if r['id']!=selected];commit(out,'Relevé supprimé.')
    else:st.info('Aucun relevé : les graphiques apparaîtront après la première saisie.')
    with st.expander('Importer des relevés CSV',key='ex_log_csv_'+p['id']):
        st.caption('Dates AAAA-MM-JJ, séparateur point-virgule, décimales virgule ou point. Le fichier doit porter exactement les colonnes du modèle ; l’import ajoute des relevés au chantier actif.')
        template=pd.DataFrame([dict(date=date.today().isoformat(),volume_m3=10.,heures_actives=2.,heures_arret=0.,carburant_l=8.,depense_fcfa=40000.,note='Exemple à remplacer')])
        st.download_button('Modèle CSV de relevés',csv_bytes(template),'terrapilot_modele_releves.csv','text/csv')
        uploaded=st.file_uploader('CSV des relevés',type=['csv'],max_upload_size=2,key='csv_'+p['id']+'_'+str(st.session_state.get('tp_csv_nonce',0)))
        if uploaded:
            try:
                parsed=import_logs(data,p,uploaded.getvalue());st.dataframe(pd.DataFrame(parsed)[LOG_COLUMNS],hide_index=True,width='stretch',alt='Aperçu des relevés CSV validés avant import')
                if st.button('Importer ces relevés',type='primary'):
                    out=copy.deepcopy(data);out['logs'].extend(parsed);st.session_state['tp_csv_nonce']=st.session_state.get('tp_csv_nonce',0)+1;commit(out,str(len(parsed))+' relevé(s) importé(s).')
            except ValueError as error:st.error(str(error))


def exports_page(data,p):
    st.header('Guide et exports')
    if p:
        st.subheader('Rapport du chantier actif');sim=cached_sim(p);fictitious=st.session_state['tp_space']=='Exemple fictif';prefix='EXEMPLE_FICTIF_' if fictitious else ''
        with st.container(horizontal=True):
            st.download_button('Rapport PDF',report_pdf(data,p,sim,fictitious),prefix+'TerraPilot_'+p['id']+'.pdf','application/pdf',icon=':material/picture_as_pdf:')
            st.download_button('Classeur Excel',report_excel(data,p,sim,fictitious),prefix+'TerraPilot_'+p['id']+'.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',icon=':material/table_view:')
            frame=pd.DataFrame(records(data,p),columns=['id','project','date','engine']+LOG_COLUMNS[1:])
            frame['origine']='Exemple fictif' if fictitious else 'Données utilisateur'
            st.download_button('Journal CSV',csv_bytes(frame),prefix+'TerraPilot_journal_'+p['id']+'.csv','text/csv',icon=':material/download:')
        st.caption('Le PDF et le classeur incluent les résultats simulés seulement si une simulation correspondant exactement à la fiche active est présente dans cette session.')
    backup=dict(application=APP,fictitious=st.session_state['tp_space']=='Exemple fictif',exported=stamp(),data=data)
    st.download_button('Sauvegarde complète JSON',json.dumps(backup,ensure_ascii=False,indent=2).encode(),'TerraPilot_sauvegarde.json','application/json')
    with st.expander('Restaurer des copies de chantiers',key='ex_restore'):
        uploaded=st.file_uploader('Sauvegarde TerraPilot',type=['json'],max_upload_size=15)
        confirmed=st.checkbox('Ajouter des copies avec de nouveaux identifiants')
        if st.button('Restaurer la sauvegarde',disabled=not uploaded or not confirmed):
            try:commit(merge_backup(data,uploaded.getvalue(),st.session_state['tp_space']=='Exemple fictif'),'Copies restaurées avec leurs relevés.')
            except ValueError as error:st.error(str(error))
    st.subheader('Un premier parcours en cinq minutes')
    st.markdown('1. Clique **Charger l’exemple fictif** dans le menu.\n2. Ouvre **Fouilles et engins** : 10 × 2 × 1,5 m = **30 m³** ; le foisonnement à 20 % donne **36 m³** à transporter.\n3. Avec la pelle compacte à 5 m³/h, le travail actif dure **6 h**. À 20 000 FCFA/h, carburant inclus, l’estimation est **120 000 FCFA**.\n4. Dans **Simulations**, choisis **Fixe** et **0 % d’arrêt** : toutes les durées valent **0,75 jour ouvré** à 8 h/j et tous les coûts **120 000 FCFA**.\n5. Dans **Suivi des travaux**, l’exemple contient 10 m³ réalisés en 2 h : **33,33 %** d’avancement et **5 m³/h** observés.')
    st.subheader('Ce que les résultats signifient')
    st.write('Le terrain et les couches sont décrits, pas diagnostiqués. La forme de la fouille sert au métré. Le rendement et le foisonnement sont des hypothèses modifiables, jamais attribuées automatiquement à un nom de sol.')
    st.write('La durée théorique n’inclut pas d’arrêt. Les simulations ajoutent une incertitude sur le rendement et un éventuel arrêt agrégé. La sensibilité compare des rangs ; elle ne prouve pas une causalité. Les jours ouvrés équivalents ne sont pas des dates calendaires.')
    st.write('Le journal décrit les interventions effectivement renseignées. Ses dépenses sont saisies indépendamment des estimations et ne sont pas recalculées après une modification des tarifs.')
    st.subheader('Conservation des données')
    st.write('Mes chantiers est enregistré dans une base SQLite située hors du dossier livré. L’exemple fictif reste séparé et n’est pas enregistré dans le registre réel. Télécharge une sauvegarde JSON pour conserver les photos, fiches, moyens et relevés ensemble.')
    st.caption('Sur un hébergement à disque éphémère, la base locale peut être perdue lors d’une reconstruction. Cette version ne promet ni stockage cloud durable, ni synchronisation hors ligne. Le README explique la préparation du déploiement.')
    st.subheader('Cadre pédagogique et sources')
    st.markdown('TerraPilot ne détermine pas une profondeur de fondation, une portance admissible ni un blindage. Pour comprendre les facteurs de danger liés aux fouilles : [INRS — fouilles en tranchée, ensevelissement](https://portaildocumentaire.inrs.fr/Default/doc/SYRACUSE/172358/les-fouilles-en-tranchee-ensevelissement?_lg=fr-FR). Cette ressource française est documentaire ; l’application n’en déduit aucune conformité locale.\n\nPour les dépendances de déploiement : [documentation officielle Streamlit](https://docs.streamlit.io/deploy/streamlit-community-cloud/deploy-your-app/app-dependencies).')


def main():
    st.set_page_config(page_title='TerraPilot · Fouilles et engins',page_icon=':material/construction:',layout='wide')
    st.html(CSS)
    st.session_state.setdefault('tp_space','Mes chantiers');st.session_state.setdefault('tp_demo',empty());st.session_state.setdefault('tp_page','Vue d’ensemble')
    try:st.session_state['tp_real'],st.session_state['tp_revision']=load_db()
    except (OSError,ValueError,sqlite3.Error) as error:st.error('Registre indisponible : '+str(error));st.stop()
    data=current()
    if data['projects'] and not active(data):st.session_state['tp_active']=data['projects'][0]['id']
    p=active(data)
    with st.sidebar:
        st.html('<div class="tp-logo"><span class="tp-symbol">T</span>TerraPilot</div><div class="tp-side-caption">FOUILLES &amp; TERRASSEMENT<br>ESTIMER · SIMULER · SUIVRE</div>')
        st.session_state['_tp_space']=st.session_state['tp_space']
        st.selectbox('Espace de travail',['Mes chantiers','Exemple fictif'],key='_tp_space',on_change=switch_space)
        st.button('Charger l’exemple fictif',on_click=load_demo,icon=':material/science:',width='stretch',wrap=True)
        st.space('small')
        for page,icon in NAV.items():st.button(page,key='nav_'+page,on_click=nav,args=(page,),icon=f':material/{icon}:',type='primary' if st.session_state['tp_page']==page else 'secondary',width='stretch',wrap=True)
        st.html('<div class="tp-side-note">CHANTIER ACTIF<br><b>'+esc(p['name'] if p else 'Aucun chantier')+'</b><br>'+esc(p['id'] if p else 'Crée ta première fiche')+'<br><br>Un atelier pédagogique de génie civil.<br>Version 1.0 · Français</div>')
    st.html('<div class="tp-topline"><span>'+esc(p['id']+' / '+p['city'] if p else 'TON ESPACE DE TRAVAIL')+'</span><span class="tp-pill'+(' demo' if st.session_state['tp_space']=='Exemple fictif' else '')+'">'+esc(st.session_state['tp_space'])+'</span></div>')
    if message:=st.session_state.pop('tp_flash',None):st.success(message)
    pages={'Vue d’ensemble':dashboard,'Chantiers':projects_page,'Fouilles et engins':geometry_page,'Simulations':simulations_page,'Suivi des travaux':tracking_page,'Guide et exports':exports_page}
    pages[st.session_state['tp_page']](data,p)
    st.html('<footer class="tp-footer">TerraPilot · Métré, comparaison de moyens et incertitudes de chantier.<br>Hypothèses saisies et résultats pédagogiques ; les relevés réels restent distincts des simulations.</footer>')


if __name__=='__main__':main()
